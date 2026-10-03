# -*- coding: utf-8 -*-
"""export_parquet.py — Parquet-Export der committeten XLSX für die R-Analyse.

Liest data/deutschland-rohdaten.xlsx (beide Blätter) und schreibt:
  - dist/deutschland-rohdaten.parquet        (400 Kreise x 109 Spalten)
  - dist/deutschland-variablen.parquet      (Registry: Bezeichnung, Einheit,
                                            Quelle, Erhebungsstand, Link,
                                            Anmerkung je Spalte)

Typsicherung für R:
  - Textspalten (rs, name, type, bl, preis_miet, name_vgrdl) als String,
    führende Nullen des Regionalschlüssels bleiben erhalten
  - numerische Spalten je Wertlage: int64 (nur Ganzzahlen) oder double
  - leere Zellen -> NA

Aufruf:  python3 src/export_parquet.py
Ausgabe: dist/deutschland-rohdaten.parquet, dist/deutschland-variablen.parquet
"""
import os
import sys

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA = os.path.join(ROOT, "data")
DIST = os.path.join(ROOT, "dist")

XLSX = os.path.join(DATA, "deutschland-rohdaten.xlsx")




def _norm(v):
    if v is None or v == "":
        return None
    return v


def load_raw(ws):
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    data = [[_norm(v) for v in r] for r in rows[1:]]
    assert len(data) == 400, len(data)
    assert len(header) == len(data[0]), (len(header), len(data[0]))
    fields = []
    cols = []
    for i, key in enumerate(header):
        vals = [r[i] for r in data]
        non_null = [v for v in vals if v is not None]
        if any(isinstance(v, str) for v in non_null):
            typ = pa.string()
        elif non_null and all(isinstance(v, int) for v in non_null):
            typ = pa.int64()
        else:
            typ = pa.float64()
            vals = [float(v) if v is not None else None for v in vals]
        fields.append(pa.field(key, typ, nullable=True))
        cols.append(vals)
    schema = pa.schema(fields)
    return pa.table({f.name: pa.array(c, type=f.type)
                     for f, c in zip(fields, cols)}, schema=schema)


def load_registry(ws):
    rows = list(ws.iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows) if r[0] == "Spalte")
    header = [str(h) for h in rows[start]]
    reg = [[_norm(v) for v in r[:len(header)]] for r in rows[start + 1:]
           if r[0] is not None]
    assert len(reg) == 109, len(reg)
    fields = [pa.field(h, pa.string(), nullable=True) for h in header]
    return pa.table({f.name: pa.array([r[i] for r in reg], type=f.type)
                     for i, f in enumerate(fields)})


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else XLSX
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    raw = load_raw(wb["Rohdaten"])
    reg = load_registry(wb["Quellen & Variablen"])

    raw_keys = set(raw.column_names)
    reg_keys = set(reg.column_names[:0])  # Spaltennamen sind Zeilen der Registry
    assert all(c in raw_keys for c in reg["Spalte"].to_pylist()), \
        "Registry-Spalten passen nicht zur Rohdaten-Tabelle"

    os.makedirs(DIST, exist_ok=True)
    out_raw = os.path.join(DIST, "deutschland-rohdaten.parquet")
    out_reg = os.path.join(DIST, "deutschland-variablen.parquet")
    pq.write_table(raw, out_raw)
    pq.write_table(reg, out_reg)
    print("OK: %s (%d Zeilen x %d Spalten)" % (out_raw, raw.num_rows, raw.num_columns))
    print("OK: %s (%d Registry-Eintraege)" % (out_reg, reg.num_rows))


if __name__ == "__main__":
    main()
